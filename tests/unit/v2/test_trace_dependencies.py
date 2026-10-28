"""S3.4 (ADR-5 «Связи показателя в модели», design D2, D3): formula walk and the ``trace_dependencies`` tool."""

import hashlib
import os
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from tests.unit.v2.fakes import FakeContext
from tests.unit.v2.test_boundaries import imported_modules

from aigw_service.api.v2.tools import dependencies as deps
from aigw_service.api.v2.tools.catalog import default_catalog
from aigw_service.api.v2.tools.context import ModelIndex
from aigw_service.api.v2.tools.dependencies import TraceArgs, cell_map, trace_dependencies
from aigw_service.api.v2.tools.envelope import ErrorCode, ToolError
from aigw_service.api.v2.tools.index_builder import build_model_index
from excel_engine import dependencies
from excel_engine.dependencies import trace

RECORD, STOP = ["Outputs"], ["Inputs"]


def workbook(**sheets: dict[str, Any]) -> Any:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, cells in sheets.items():
        sheet = wb.create_sheet(title.replace("_", " "))
        for coordinate, value in cells.items():
            sheet[coordinate] = value
    return wb


class TestTrace:
    """Depth, ranges, cross-sheet references, cycles, the cell limit and references the walk cannot follow."""

    def test_inputs_with_smallest_depth(self) -> None:
        wb = workbook(
            Inputs={"B2": 10, "B3": 20},
            Calc={"A1": "=Inputs!B2*2", "A2": "=Inputs!B3+A1+Inputs!B2"},
            Outputs={"C2": "=Calc!A2"},
        )
        rows = trace(wb, ("Outputs", "C2"), RECORD, STOP).rows
        assert rows[("Inputs", 3)] == 2
        assert rows[("Inputs", 2)] == 2  # also reachable at depth 3 through A1, the smallest depth wins

    def test_range_is_expanded(self) -> None:
        wb = workbook(Inputs={"B2": 1, "B3": 2, "B4": 3}, Outputs={"C2": "=SUM(Inputs!$B$2:$B$4)"})
        rows = trace(wb, ("Outputs", "C2"), RECORD, STOP).rows
        assert {r for (s, r) in rows if s == "Inputs"} == {2, 3, 4}

    def test_quoted_sheet_and_intermediate_output(self) -> None:
        wb = workbook(
            Inputs={"B2": 1},
            P_L={"A1": "=Inputs!B2"},
            Outputs={"C2": "='P L'!A1*2", "C3": "=C2/4"},
        )
        rows = trace(wb, ("Outputs", "C3"), RECORD, STOP).rows
        assert rows[("Outputs", 2)] == 1
        assert rows[("Inputs", 2)] == 3

    def test_cycle_and_repeated_paths_end(self) -> None:
        wb = workbook(
            Inputs={"B2": 1},
            Calc={"A1": "=A2+Inputs!B2", "A2": "=A1+Inputs!B2"},
            Outputs={"C2": "=Calc!A1+Calc!A2"},
        )
        result = trace(wb, ("Outputs", "C2"), RECORD, STOP)
        assert result.rows[("Inputs", 2)] == 2
        assert result.visited == 4
        assert not result.truncated

    def test_input_cells_are_leaves(self) -> None:
        wb = workbook(Inputs={"B2": "=B3*2", "B3": 5}, Outputs={"C2": "=Inputs!B2"})
        rows = trace(wb, ("Outputs", "C2"), RECORD, STOP).rows
        assert ("Inputs", 3) not in rows

    def test_cut_at_cell_limit(self) -> None:
        wb = workbook(Inputs={"B2": 1}, Calc={f"A{i}": f"=A{i + 1}" for i in range(1, 50)}, Outputs={"C2": "=Calc!A1"})
        wb["Calc"]["A50"] = "=Inputs!B2"
        result = trace(wb, ("Outputs", "C2"), RECORD, STOP, max_cells=10)
        assert result.truncated
        assert result.visited == 10
        assert ("Inputs", 2) not in result.rows

    @pytest.mark.parametrize(
        "formula", ["=[1]Inputs!B2", "=Revenue*2", "=SUM(Inputs!B:B)", "=Missing!A1", "=SUM(Inputs!A1:Z1000)"]
    )
    def test_reference_not_followed_is_counted(self, formula: str) -> None:
        wb = workbook(Inputs={"B2": 1}, Outputs={"C2": formula})
        result = trace(wb, ("Outputs", "C2"), RECORD, STOP, max_cells=100)
        assert result.skipped == 1
        assert ("Inputs", 2) not in result.rows

    def test_empty_cells_are_not_created(self) -> None:
        """The cached workbook is shared between calls: reading a range must not add cells to it."""
        wb = workbook(Inputs={"B2": 1}, Outputs={"C2": "=SUM(Inputs!B2:B9)"})
        trace(wb, ("Outputs", "C2"), RECORD, STOP)
        assert wb["Inputs"].max_row == 2

    def test_unknown_start_sheet(self) -> None:
        with pytest.raises(KeyError):
            trace(workbook(Inputs={}), ("Outputs", "C2"), RECORD, STOP)


def test_walk_imports_neither_framework_nor_v1() -> None:
    """AR-Q3: the engine copy stays framework-free and does not reach into v1."""
    modules = imported_modules(Path(dependencies.__file__))
    assert not [m for m in modules if m.startswith(("langchain", "langgraph", "aigw_service"))]


# ---------------------------------------------------------------- the tool on a small model workbook

MODEL_INDEX = ModelIndex(
    {"IN0": "Цена метанола", "IN1": "Курс USD", "IN2": "Капекс"},
    {"OUT0": "Выручка", "OUT1": "Раздел: показатели", "OUT2": "EBITDA"},
    output_years={"OUT0": (2026, 2027), "OUT1": (), "OUT2": (2026, 2027)},
)


def model_workbook(path: Path) -> str:
    """Inputs → Calc → Outputs; EBITDA goes through Выручка; Капекс is not connected."""
    wb = workbook(
        Inputs={
            "A1": "Наименование",
            "B1": 2026,
            "C1": 2027,
            "A2": "Цена метанола",
            "B2": 400,
            "C2": 410,
            "A3": "Курс USD",
            "B3": 90,
            "C3": 92,
            "A4": "Капекс",
            "B4": 50,
            "C4": 50,
        },
        Calc={"B5": "=Inputs!B2*Inputs!B3", "C5": "=Inputs!C2*Inputs!C3"},
        Outputs={
            "A1": "Наименование",
            "B1": 2026,
            "C1": 2027,
            "A2": "Выручка",
            "B2": "=Calc!B5",
            "C2": "=Calc!C5",
            "A3": "Раздел: показатели",
            "A4": "EBITDA",
            "B4": "=B2-100",
            "C4": "=C2-100",
        },
    )
    target = path / "model.xlsx"
    wb.save(target)
    return str(target)


def context(path: str) -> FakeContext:
    return FakeContext(file_path=path, model_index=MODEL_INDEX)


class TestCellMap:
    def test_numbering_matches_model_index(self, tmp_path: Path) -> None:
        path = model_workbook(tmp_path)
        cells, index = cell_map(path), build_model_index(path)
        assert {i: index.inputs[i] for i in cells.input_ids.values()} == dict(index.inputs)
        assert set(cells.output_rows) == set(index.outputs)
        assert cells.output_rows["OUT2"] == 4
        assert cells.input_ids[3] == "IN1"
        assert cells.year_columns == {2026: "B", 2027: "C"}

    def test_cached_until_the_file_changes(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        path = model_workbook(tmp_path)
        loads: list[bool] = []
        real = openpyxl.load_workbook
        monkeypatch.setattr(
            deps.openpyxl, "load_workbook", lambda p, **kw: loads.append(kw["data_only"]) or real(p, **kw)
        )
        deps._load.cache_clear()  # pylint: disable=protected-access
        first = deps.load(path)
        assert deps.load(path) is first
        assert len(loads) == 2
        stat = os.stat(path)
        os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000))
        assert deps.load(path) is not first
        assert len(loads) == 4


class TestTool:
    def test_inputs_and_intermediate_output(self, tmp_path: Path) -> None:
        result = trace_dependencies(TraceArgs(output_id="OUT2", year=2026), context(model_workbook(tmp_path)))
        assert [(i["id"], i["name"], i["depth"]) for i in result.data["inputs"]] == [
            ("IN0", "Цена метанола", 3),
            ("IN1", "Курс USD", 3),
        ]
        assert result.data["outputs"] == [{"id": "OUT0", "name": "Выручка", "depth": 1}]
        assert "Цена метанола" in result.result
        assert "Выручка" in result.result
        assert not result.notes
        assert result.hint
        assert [n.value for n in result.numbers] == [2, result.data["visited"]]

    def test_repeated_input_name_is_reported(self, tmp_path: Path) -> None:
        """Found in S3.4 on Methanex (Henry Hub on 3 rows): the index keeps the last row, the formulas use another."""
        path = model_workbook(tmp_path)
        wb = openpyxl.load_workbook(path)
        wb["Inputs"]["A5"], wb["Inputs"]["B5"] = "Курс USD", 95  # same name; the index now points at row 5
        wb.save(path)
        cells = cell_map(path)
        assert cells.input_ids[3] == cells.input_ids[5] == "IN1"
        assert cells.input_rows["IN1"] == 5
        result = trace_dependencies(TraceArgs(output_id="OUT2", year=2026), context(path))
        assert "IN1" in [i["id"] for i in result.data["inputs"]]
        assert any("«Курс USD»" in n and "повторяется" in n for n in result.notes)

    def test_truncated_walk_is_reported(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        monkeypatch.setattr(deps, "trace", lambda *a, **_kw: trace(*a, max_cells=2))
        result = trace_dependencies(TraceArgs(output_id="OUT2", year=2027), context(model_workbook(tmp_path)))
        assert result.data["truncated"]
        assert deps.TRUNCATED_NOTE in result.notes

    @pytest.mark.parametrize(
        ("args", "code"),
        [
            (TraceArgs(year=2026), ErrorCode.MISSING_INPUT),
            (TraceArgs(output_id="OUT2"), ErrorCode.MISSING_INPUT),
            (TraceArgs(output_id="OUT9", year=2026), ErrorCode.UNKNOWN_ID),
            (TraceArgs(output_id="OUT1", year=2026), ErrorCode.NO_VALUES),
            (TraceArgs(output_id="OUT2", year=2030), ErrorCode.INVALID_ARGS),
        ],
    )
    def test_argument_errors_before_opening(
        self, args: TraceArgs, code: ErrorCode, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        monkeypatch.setattr(deps, "load", lambda _p: pytest.fail("the workbook must not be opened"))
        with pytest.raises(ToolError) as caught:
            trace_dependencies(args, context("/nonexistent.xlsx"))
        assert caught.value.code is code
        assert caught.value.hint

    def test_workbook_file_unchanged(self, tmp_path: Path) -> None:
        path = model_workbook(tmp_path)
        before = (os.stat(path).st_mtime_ns, hashlib.sha256(Path(path).read_bytes()).hexdigest())
        trace_dependencies(TraceArgs(output_id="OUT2", year=2026), context(path))
        assert (os.stat(path).st_mtime_ns, hashlib.sha256(Path(path).read_bytes()).hexdigest()) == before


class TestCatalog:
    """ADR-5 «Каталог инструментов»: trace_dependencies is opened by sensitivity-drivers and needs no engine."""

    def test_registered_as_skill_tool(self) -> None:
        catalog = default_catalog()
        assert "trace_dependencies" in {t.name for t in catalog.skill_tools()}
        assert "trace_dependencies" not in {t.name for t in catalog.base_tools()}
        assert "trace_dependencies" not in catalog.engine_tool_names()
        assert catalog.tool_skills() == {"trace_dependencies": "sensitivity-drivers"}

    def test_schema_hides_service_arguments(self) -> None:
        schema = default_catalog().tool("trace_dependencies").tool_call_schema.model_json_schema()
        assert set(schema["properties"]) == {"output_id", "year"}
