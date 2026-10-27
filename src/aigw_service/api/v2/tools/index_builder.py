"""Model index from the uploaded workbook: IDs, names, values by year, units and output types (spec §2 «Индекс модели»).

One openpyxl pass (cached values, ``data_only=True``); rows go to the same v1 mapping functions as ``build_mappings``,
so IDs keep the ``_build_catalog_with_ids`` numbering. Units and types are read from the same rows by header (S1.7 D4).
"""

from typing import Any

import openpyxl

from aigw_service.api.v1.tools import create_input_mapping, create_output_mapping
from aigw_service.api.v2.tools.context import ModelIndex

UNIT_HEADER = "Ед. измерения"
TYPE_HEADER = "Тип показателя"
NO_UNIT = {"", "[требуется заполнить]"}


def _numeric(values: dict) -> dict[int, float]:
    return {int(year): float(v) for year, v in (values or {}).items() if isinstance(v, (int, float))}


def _sheet_rows(workbook: Any, name: str) -> list[list[Any]]:
    sheet = workbook[name]
    return [
        [c.value for c in row] for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column)
    ]


def _cell_text(rows: list[list[Any]], row: int, header: str) -> str | None:
    """Text of the ``header`` column in 1-based ``row``; ``None`` for a missing column or a placeholder."""
    headers = rows[0] if rows else []
    if header not in headers or row - 1 >= len(rows):
        return None
    line = rows[row - 1]
    column = headers.index(header)
    text = str(line[column]).strip() if column < len(line) and line[column] is not None else ""
    return None if text in NO_UNIT else text


def build_model_index(path: str) -> ModelIndex:
    """One openpyxl pass; IDs are numbered in the same order as ``_build_catalog_with_ids`` (IN0…, OUT0…)."""
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        input_rows, output_rows = _sheet_rows(workbook, "Inputs"), _sheet_rows(workbook, "Outputs")
    finally:
        workbook.close()
    input_mapping, output_mapping = create_input_mapping(input_rows), create_output_mapping(output_rows)
    inputs, input_values, units = {}, {}, {}
    for idx, info in enumerate(input_mapping["row_mapping"].values()):
        key = f"IN{idx}"
        inputs[key] = info["original"].strip()
        input_values[key] = _numeric(info.get("values"))
        if unit := _cell_text(input_rows, info["row"], UNIT_HEADER):
            units[key] = unit
    outputs, output_years, output_values, output_types = {}, {}, {}, {}
    for idx, info in enumerate(output_mapping["output_mapping"].values()):
        key = f"OUT{idx}"
        outputs[key] = info["original"].strip()
        output_values[key] = _numeric(info.get("values"))
        output_years[key] = tuple(sorted(output_values[key]))
        if unit := _cell_text(output_rows, info["row"], UNIT_HEADER):
            units[key] = unit
        if kind := _cell_text(output_rows, info["row"], TYPE_HEADER):
            output_types[key] = kind
    return ModelIndex(inputs, outputs, input_values, output_years, output_values, units, output_types)
