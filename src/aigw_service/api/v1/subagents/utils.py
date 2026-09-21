from typing import Any

import openpyxl

from aigw_service.api.v1.tools import (
    create_input_mapping,
    create_output_mapping,
)


def save_graph_image(graph):
    save_path = "graph.png"
    with open(save_path, "wb") as f:
        f.write(graph.get_graph(xray=True).draw_mermaid_png())


def _build_catalog_with_ids(model_path: str) -> tuple[dict[str, str], dict[str, str]]:
    input_mapping, output_mapping = build_mappings(model_path)
    inputs_dict = {}
    outputs_dict = {}
    for idx, info in enumerate(input_mapping["row_mapping"].values()):
        inputs_dict[f"IN{idx}"] = info["original"].strip()

    for idx, info in enumerate(output_mapping["output_mapping"].values()):
        outputs_dict[f"OUT{idx}"] = info["original"].strip()
    return inputs_dict, outputs_dict


def _build_catalog(model_path: str) -> tuple[str, str]:
    input_mapping, output_mapping = build_mappings(model_path)
    inputs = sorted({info["original"].strip() for info in input_mapping["row_mapping"].values()}, key=str.lower)
    outputs = sorted({info["original"].strip() for info in output_mapping["output_mapping"].values()}, key=str.lower)
    return "\n".join(inputs), "\n".join(outputs)


def build_mappings(model_path: str) -> tuple[dict, dict]:
    wb = openpyxl.load_workbook(model_path, data_only=True)
    inputs_data = [
        [c.value for c in row]
        for row in wb["Inputs"].iter_rows(min_row=1, max_row=wb["Inputs"].max_row, max_col=wb["Inputs"].max_column)
    ]
    outputs_data = [
        [c.value for c in row]
        for row in wb["Outputs"].iter_rows(min_row=1, max_row=wb["Outputs"].max_row, max_col=wb["Outputs"].max_column)
    ]
    wb.close()
    return create_input_mapping(inputs_data), create_output_mapping(outputs_data)


def _stringify_catalog(data: dict[str, str]) -> str:
    output_str = ""
    for key, value in data.items():
        output_str += f"ID: {key} | Название: {value}\n"
    return output_str


def get_tokens(response: Any) -> dict[str, int]:
    """Извлекает token_usage из ответа модели (порядок fallback'ов определяет приоритет)."""
    if response is None:
        return {}

    # 1) response_metadata["token_usage"] — для structured_output(include_raw=True) и GigaChat
    rm = getattr(response, "response_metadata", None) or {}
    tu = rm.get("token_usage", {})
    if tu.get("total_tokens", 0) > 0:
        return tu

    # 2) llm_output["token_usage"] — GigaChat
    lo = getattr(response, "llm_output", None) or {}
    tu = lo.get("token_usage", {})
    if tu.get("total_tokens", 0) > 0:
        return tu

    # 3) Standard UsageMetadata (работает и для GigaChat, и для Ollama)
    um = getattr(response, "usage_metadata", None)
    if um:
        return {
            "prompt_tokens": um.input_tokens,
            "completion_tokens": um.output_tokens,
            "total_tokens": um.total_tokens,
        }

    return {}


def build_diagnostic_data(
    node: str,
    elapsed: float,
    token_usage: dict[str, int],
    extra: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Формирует диагностический словарь для логирования LLM-вызовов.

    Параметры:
        node      — имя узла графа (для метрики ``node``).
        elapsed   — время выполнения в секундах (округляется до 3 знаков).
        token_usage — словарь с информацией о токенах.
        extra     — произвольные дополнительные поля (например ``{"next_agent": ...}`).
    """
    data: dict[str, Any] = {
        "node": node,
        "elapsed_s": round(elapsed, 3),
        "token_usage": token_usage,
    }
    if extra:
        data.update(extra)
    return data
