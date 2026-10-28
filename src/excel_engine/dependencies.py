"""Formula dependencies of one cell (story S3.4, design D2): breadth-first walk over an openpyxl workbook.

Based on the v1 ``get_direct_dependencies`` / ``parse_cell_ref`` (``openpyxl`` tokenizer, ``RANGE`` operands); v1 bodies
are not changed. Differences: ranges are expanded into cells, every cell keeps its smallest depth, cells of the stop
sheets (the model inputs) are leaves, and the walk stops after ``max_cells`` cells. References the walk cannot follow
(other workbooks, defined names, whole columns or rows, unknown sheets, oversized ranges) are counted.
"""

from collections import deque
from collections.abc import Collection
from dataclasses import dataclass, field
from typing import Any

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter, range_boundaries
from openpyxl.utils.exceptions import CellCoordinatesException

MAX_TRACE_CELLS = 20000
"""Cells one walk may visit; a larger graph is cut and reported as ``truncated``."""

Cell = tuple[str, str]
"""``(sheet title, coordinate)`` with the sheet title as the workbook spells it and no ``$`` in the coordinate."""


@dataclass
class Trace:
    """What the walk reached: ``rows[(sheet, row)]`` is the smallest depth of a cell of a recorded sheet."""

    rows: dict[tuple[str, int], int] = field(default_factory=dict)
    visited: int = 0
    truncated: bool = False
    skipped: int = 0


def _existing_cell(sheet: Any, coordinate: str) -> Any:
    """The cell if the sheet has it; ``sheet[coordinate]`` would create an empty one in a shared, cached workbook."""
    column, row = coordinate_from_string(coordinate)
    return sheet._cells.get((row, column_index_from_string(column)))  # pylint: disable=protected-access


def _formula(cell: Any) -> str | None:
    if cell is None:
        return None
    value = getattr(cell.value, "text", cell.value)  # ArrayFormula keeps its text in ``text``
    if cell.data_type != "f" or not isinstance(value, str):
        return None
    return value if value.startswith("=") else f"={value}"


def _split_sheet(ref: str, current: str) -> tuple[str, str]:
    if "!" not in ref:
        return current, ref
    sheet, address = ref.rsplit("!", 1)
    if sheet.startswith("'") and sheet.endswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    return sheet, address


def _cells(address: str, budget: int) -> list[str] | None:
    """Coordinates of ``A1`` or ``A1:B10``; ``None`` for a defined name, a whole column / row or a range over budget."""
    address = address.replace("$", "")
    try:
        if ":" not in address:
            column, row = coordinate_from_string(address)
            return [f"{column}{row}"]
        min_col, min_row, max_col, max_row = range_boundaries(address)
    except (ValueError, CellCoordinatesException):  # a defined name, not a cell address
        return None
    if None in (min_col, min_row, max_col, max_row):
        return None  # whole column or row
    size = (max_col - min_col + 1) * (max_row - min_row + 1)  # type: ignore[operator]
    if size > budget:
        return None
    return [
        f"{get_column_letter(c)}{r}"
        for r in range(min_row, max_row + 1)  # type: ignore[arg-type]
        for c in range(min_col, max_col + 1)  # type: ignore[arg-type]
    ]


class _Walk:  # pylint: disable=too-few-public-methods,too-many-instance-attributes  # state of one walk
    def __init__(self, workbook: Any, record: Collection[str], stop: Collection[str], max_cells: int) -> None:
        self.workbook = workbook
        self.titles = {name.lower(): name for name in workbook.sheetnames}
        self.record = {s.lower() for s in record} | {s.lower() for s in stop}
        self.stop = {s.lower() for s in stop}
        self.max_cells = max_cells
        self.trace = Trace()
        self.seen: set[Cell] = set()
        self.queue: deque[tuple[Cell, int]] = deque()

    def push(self, cell: Cell, depth: int) -> None:
        if cell not in self.seen:
            self.seen.add(cell)
            self.queue.append((cell, depth))

    def references(self, formula: str, sheet: str) -> list[Cell]:
        try:
            tokens = Tokenizer(formula).items
        except Exception:  # pylint: disable=broad-exception-caught  # an unparsable formula is a skipped reference
            self.trace.skipped += 1
            return []
        found: list[Cell] = []
        for token in tokens:
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            ref_sheet, address = _split_sheet(token.value, sheet)
            title = self.titles.get(ref_sheet.lower())
            cells = _cells(address, self.max_cells) if title and "[" not in ref_sheet else None
            if title is None or cells is None:
                self.trace.skipped += 1
                continue
            found += [(title, c) for c in cells]
        return found

    def run(self, start: Cell) -> Trace:
        title = self.titles.get(start[0].lower())
        if title is None:
            raise KeyError(f"no sheet {start[0]!r} in the workbook")
        self.push((title, start[1].replace("$", "")), 0)
        while self.queue:
            if self.trace.visited >= self.max_cells:
                self.trace.truncated = True
                break
            (sheet, coordinate), depth = self.queue.popleft()
            self.trace.visited += 1
            key = sheet.lower()
            if key in self.record:
                row = coordinate_from_string(coordinate)[1]
                self.trace.rows.setdefault((sheet, row), depth)  # breadth-first: the first depth is the smallest
            if key in self.stop and depth > 0:
                continue
            formula = _formula(_existing_cell(self.workbook[sheet], coordinate))
            for ref in self.references(formula, sheet) if formula else []:
                self.push(ref, depth + 1)
        return self.trace


def trace(
    workbook: Any,
    start: Cell,
    record_sheets: Collection[str],
    stop_sheets: Collection[str],
    max_cells: int = MAX_TRACE_CELLS,
) -> Trace:
    """Walk the formulas from ``start``; rows of ``record_sheets`` and ``stop_sheets`` are recorded with their depth.

    ``workbook`` is an openpyxl workbook loaded with formulas (``data_only=False``); it is only read.
    """
    return _Walk(workbook, record_sheets, stop_sheets, max_cells).run(start)


def column_letter(index: int) -> str:
    """Excel letter of a 0-based column index (the year columns of the mapping functions are 0-based)."""
    return get_column_letter(index + 1)
